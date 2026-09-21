from docx import Document
from pathlib import Path
p=Path('tests/output');p.mkdir(parents=True,exist_ok=True)
d=Document();d.add_heading('第一章 緒論',1);d.add_paragraph('Test results: t(20) = 2.5, p = .001. p = 1.4. SD = -2. 95% CI [5, 2].');d.add_paragraph('This text and its citation (Author, 2024) must remain unchanged.');d.add_heading('參考文獻',1);d.add_paragraph('Author (2024). Research title. https://doi.org/10.1038/nphys1170');d.add_paragraph('Unknown (2020). Missing research. https://doi.org/10.9999/not-a-real-record');d.save(p/'fixture.docx')
